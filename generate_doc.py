# -*- coding: utf-8 -*-
"""Generate PromptClip usage guide as Word document."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)
OUTPUT_PATH = os.path.join(OUTPUT_DIR, "PromptClip-使用说明.docx")

doc = Document()

# Page setup
for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles['Normal']
font = style.font
font.name = 'Microsoft YaHei'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def set_font(run, name='Microsoft YaHei', size=None, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_page_break():
    doc.add_page_break()


def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    return h


def add_para(text, bold=False, size=11, alignment=None, color=None, space_after=6):
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run(text)
    set_font(run, size=11)
    return p


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        set_font(run, size=10, bold=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = ''
            run = cell.paragraphs[0].add_run(val)
            set_font(run, size=10)
    doc.add_paragraph()


# ===== FOOTER =====
for section in doc.sections:
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("PromptClip | 提示词剪刀 使用说明")
    set_font(run, size=9, color=(128, 128, 128))


# ==========================================
# COVER PAGE
# ==========================================
for _ in range(6):
    doc.add_paragraph()

add_para("PromptClip | 提示词剪刀", bold=True, size=32,
         alignment=WD_ALIGN_PARAGRAPH.CENTER, color=(139, 26, 26))

add_para("使用说明", bold=True, size=24,
         alignment=WD_ALIGN_PARAGRAPH.CENTER, color=(139, 26, 26))

doc.add_paragraph()
doc.add_paragraph()

add_para("面向 6-12 岁低龄儿童的 AI 提示词拼装工具", bold=False, size=14,
         alignment=WD_ALIGN_PARAGRAPH.CENTER, color=(100, 100, 100))

doc.add_paragraph()
doc.add_paragraph()

add_para("适用对象：教师、学生", bold=False, size=12,
         alignment=WD_ALIGN_PARAGRAPH.CENTER, color=(80, 80, 80))
add_para("版本：v1.0", bold=False, size=12,
         alignment=WD_ALIGN_PARAGRAPH.CENTER, color=(80, 80, 80))
add_para("日期：2026 年 5 月", bold=False, size=12,
         alignment=WD_ALIGN_PARAGRAPH.CENTER, color=(80, 80, 80))

doc.add_paragraph()
doc.add_paragraph()

add_para("在线访问：https://promptclip-pi.vercel.app", bold=False, size=11,
         alignment=WD_ALIGN_PARAGRAPH.CENTER, color=(100, 100, 100))

# ==========================================
# PAGE BREAK -> TOC
# ==========================================
add_page_break()

add_heading_styled("目  录", level=1)
doc.add_paragraph()

toc_items = [
    ("一、这是什么？", "3"),
    ("二、学生篇：怎么用？", "4"),
    ("    第一步：打开页面", "4"),
    ("    第二步：选择卡片", "4"),
    ("    第三步：复制提示词", "5"),
    ("    第四步：粘贴到 AI 绘画工具", "5"),
    ("    其他按钮", "5"),
    ("三、老师篇：怎么管理？", "6"),
    ("    打开老师模式 / 新增词条 / 删除词条", "6"),
    ("    导出 / 导入词库 / 课堂模板", "7"),
    ("四、课堂使用建议", "8"),
    ("五、常见问题", "9"),
    ("六、技术信息", "10"),
]

for item, page in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    dots = "." * (50 - len(item))
    run = p.add_run("{}{} {}".format(item, dots, page))
    set_font(run, size=11)

# ==========================================
# PAGE BREAK -> CHAPTER 1
# ==========================================
add_page_break()

add_heading_styled("一、这是什么？", level=1)

add_para(
    "PromptClip（提示词剪刀）是一个面向 6-12 岁低龄儿童的提示词拼装工具。"
    "孩子不用打字，通过点击大号卡片，就能像搭积木一样完成"
    "“画什么 + 画成什么样 + 细节”的提示词组合，"
    "然后一键复制给 AI 绘画工具使用。"
)
add_para(
    "支持的 AI 绘画工具包括：通义万相、文心一格、Midjourney、"
    "Stable Diffusion 网页版等。"
)

# ==========================================
# CHAPTER 2: Student
# ==========================================
add_page_break()
add_heading_styled("二、学生篇：怎么用？", level=1)

# Step 1
add_heading_styled("第一步：打开页面", level=2)
add_para("用浏览器打开网站，你会看到顶部有一个红色标题栏，写着：")
add_para("提示词 = 画什么 + 画成什么样 + 细节", bold=True, size=12,
         alignment=WD_ALIGN_PARAGRAPH.CENTER)

# Step 2
add_heading_styled("第二步：选择卡片", level=2)
add_para("页面主体分成三列，用不同颜色区分：")

add_table(
    ["列", "颜色", "选几个", "说明"],
    [
        ["① 主体（画什么）", "粉红色", "选 1 个",
         "选择你想画的主角，如“雪豹”“小牦牛”"],
        ["② 风格（画成什么样）", "蓝色", "选 1 个",
         "选择画风，如“可爱卡通”“水彩”"],
        ["③ 细节（还有什么）", "绿色", "可多选",
         "添加画面元素，如“五色经幡”“远处雪山”"],
    ]
)

add_bullet("点击卡片 → 选中（卡片会变色变大）")
add_bullet("再点一下 → 取消选择")

# Step 3
add_heading_styled("第三步：复制提示词", level=2)
add_para("页面底部会自动显示拼好的提示词，例如：")
add_para(
    "请画雪豹，使用可爱卡通风格，画面中加入五色经幡、"
    "远处雪山、1:1 方形构图，整体温暖明亮，适合低龄儿童 AI 创作。",
    bold=False, size=11
)
add_para(
    "点击“复制给 AI”按钮，看到绿色提示框显示"
    "“已复制，可以粘贴给 AI 了！”，就说明复制成功了。"
)

# Step 4
add_heading_styled("第四步：粘贴到 AI 绘画工具", level=2)
add_para("打开你常用的 AI 绘画工具，在输入框中粘贴（Ctrl+V），发送即可。")

# Other buttons
add_heading_styled("其他按钮", level=2)
add_table(
    ["按钮", "作用"],
    [
        ["换一组灵感", "随机帮你选一组搭配，看看能出现什么惊喜"],
        ["清空选择", "取消所有已选的卡片，重新开始"],
        ["课堂模板", "点击模板按钮，快速套用老师预设的主题"],
    ]
)

# ==========================================
# CHAPTER 3: Teacher
# ==========================================
add_page_break()
add_heading_styled("三、老师篇：怎么管理？", level=1)

add_heading_styled("打开老师模式", level=2)
add_para(
    "点击右上角的“老师模式：关”按钮，"
    "切换为“老师模式：开”。"
    "此时页面底部会出现老师管理面板，包含以下功能。"
)

add_heading_styled("新增词条", level=2)
add_bullet("在下拉菜单中选择要添加到哪个分类（主体 / 风格 / 细节）")
add_bullet("在输入框中输入新词条，比如“白唇鹿”")
add_bullet("点击“添加”按钮")
add_bullet("新增的词条会出现在对应的卡片列中，右上角有标记")

add_heading_styled("删除自定义词条", level=2)
add_para(
    "老师模式下，鼠标移到自定义词条（带标记）的卡片上，"
    "左上角会出现 × 按钮，点击即可删除。"
)
add_para("注意：默认词条不可删除，只有你自己添加的词条才能删除。", bold=True)

add_heading_styled("导出 / 导入词库", level=2)
add_bullet("导出词库 JSON：把当前词库（包括你添加的词条）下载为一个 JSON 文件，可以分享给其他老师，或作为备份")
add_bullet("导入词库 JSON：上传之前导出的 JSON 文件，恢复词库")
add_bullet("恢复默认词库：清空所有自定义内容，恢复到系统最初的内置词库")

add_heading_styled("课堂模板", level=2)
add_para("页面顶部有 4 个内置模板按钮：")
add_table(
    ["模板名称", "预设内容"],
    [
        ["高原动物徽章", "雪豹 + 可爱卡通 + 五色经幡、雪山、方形构图"],
        ["石渠文创明信片", "扎溪卡草原 + 水彩 + 藏八宝纹样、A6 尺寸"],
        ["藏文化角色设计", "格萨尔王 + 藏式插画 + 藏戏面具、背景虚化"],
        ["星空草原插画", "星空银河 + 梦幻童话 + 草原、经幡、雪山"],
    ]
)
add_para("点击任意模板按钮，自动帮学生填好对应选择。")

# ==========================================
# CHAPTER 4: Classroom
# ==========================================
add_page_break()
add_heading_styled("四、课堂使用建议", level=1)

add_heading_styled("场景一：自由创作（10-15 分钟）", level=2)
add_bullet("打开网页，让学生自由探索三列卡片")
add_bullet("每人组合一个自己喜欢的提示词")
add_bullet("复制后粘贴到 AI 绘画工具中生成图片")
add_bullet("全班展示，看看谁的画面最有趣")

add_heading_styled("场景二：主题课堂（20-30 分钟）", level=2)
add_bullet("老师提前通过老师模式添加本节课的主题词条")
add_bullet("导出词库 JSON 提前分享给其他班级的老师")
add_bullet("课上让学生使用指定模板")
add_bullet("学生可以在模板基础上修改细节")
add_bullet("生成后讨论：为什么选择这些词？画面和想象的一样吗？")

add_heading_styled("场景三：跨学科融合", level=2)
add_table(
    ["学科", "活动建议"],
    [
        ["语文课", "用生成的图片编故事"],
        ["美术课", "对比 AI 生成的画与手绘的区别"],
        ["自然课", "选择高原动物，查阅资料了解真实样貌"],
        ["地方文化课", "使用藏文化相关词条，了解非遗元素"],
    ]
)

# ==========================================
# CHAPTER 5: FAQ
# ==========================================
add_heading_styled("五、常见问题", level=1)

faqs = [
    ("学生需要注册账号吗？",
     "不需要。打开网页就能用，没有任何注册或登录步骤。"),
    ("需要安装 App 吗？",
     "不需要。电脑、平板浏览器直接打开网址即可。"),
    ("复制后在哪里粘贴？",
     "打开 AI 绘画工具（如通义万相、文心一格、Midjourney、"
     "Stable Diffusion 网页版等），在输入框中按 Ctrl+V 粘贴即可。"),
    ("老师添加的词条会一直保留吗？",
     "词条保存在浏览器的 localStorage 中。如果换了一台设备或清除了浏览器数据，"
     "需要重新导入之前导出的 JSON 文件。建议定期导出备份。"),
    ("可以用手机吗？",
     "页面适配了平板和电脑屏幕。手机上也能打开，但屏幕较小，"
     "建议使用平板或电脑获得更好的操作体验。"),
]

for q, a in faqs:
    add_para("Q：{}".format(q), bold=True, size=11)
    add_para("A：{}".format(a), size=11, space_after=10)

# ==========================================
# CHAPTER 6: Tech
# ==========================================
add_heading_styled("六、技术信息", level=1)
add_bullet("纯前端应用，无需后端服务器，无用户数据上传")
add_bullet("数据保存在浏览器本地（localStorage），不上传任何个人信息")
add_bullet("支持离线使用（首次加载后浏览器会缓存资源）")
add_bullet("技术栈：Vite + React + TypeScript")
add_bullet("在线地址：https://promptclip-pi.vercel.app")

# ===== SAVE =====
doc.save(OUTPUT_PATH)
print("OK: {}".format(OUTPUT_PATH))
print("Size: {} bytes".format(os.path.getsize(OUTPUT_PATH)))
